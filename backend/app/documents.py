"""Document pipeline for "chat with your files" (RAG).

Flow for one uploaded file:
    1. Parse it into text, page by page (PDF, DOCX, TXT, Markdown). Scanned PDF
       pages and image files are read with OCR first.
    2. Split the text into overlapping chunks of about 1,000 characters.
    3. Embed every chunk with FastEmbed, a small model that runs locally on the CPU.
    4. Store the chunks and their vectors in Chroma, a vector database kept on disk.
    5. Also store the chunk text in a SQLite FTS5 table, a keyword index on disk.

At question time, `search` runs two searches over the documents attached to the
chat: Chroma for chunks close in meaning, and the keyword index (BM25) for chunks
sharing exact words such as numbers, names or error codes. The two ranked lists
are merged with reciprocal rank fusion.

Everything here is free and local (including OCR): no API key is needed and the
files never leave this machine. Only the few retrieved passages are sent to the chat model.
"""

import hashlib
import io
import json
import os
import re
import sqlite3
import threading
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

from langchain_text_splitters import RecursiveCharacterTextSplitter

from . import web

# Where the vector index, the document list and the downloaded model are kept.
DATA_DIR = Path(os.getenv("WORLDBEE_DATA_DIR", Path(__file__).resolve().parent.parent / "data"))

# Any model FastEmbed supports works here. Changing it means re-uploading your
# files, because vectors from different models can't be compared.
EMBEDDING_MODEL = os.getenv("WORLDBEE_EMBEDDING_MODEL", "BAAI/bge-small-en-v1.5")

IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
SUPPORTED_TYPES = {".pdf", ".docx", ".txt", ".md"} | IMAGE_TYPES
MAX_FILE_BYTES = 20 * 1024 * 1024
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150
# How many chunks one search returns to the agent.
TOP_K = 6
# How many candidates each side (vector and keyword) brings to the merge.
CANDIDATES = 20
# Reciprocal rank fusion constant: a chunk at rank r scores 1 / (RRF_K + r) per list.
RRF_K = 60
# A PDF page with less text than this is treated as a scan and read with OCR.
MIN_PAGE_CHARS = 20
# OCR takes a second or two per page on a CPU, so only this many scanned pages are read.
MAX_OCR_PAGES = 30
OCR_SCALE = 2.5  # PDF pages are rendered at 2.5x (about 180 dpi) before OCR.
OCR_MIN_CONFIDENCE = 0.5
OCR_MAX_SIDE = 3000  # Bigger images are shrunk first, which keeps OCR fast.


class DocumentError(ValueError):
    """A problem with an uploaded file that the user can fix (wrong type, no text...)."""


# ---------- Parsing ----------


@dataclass
class Parsed:
    """The text of a file, page by page."""

    pages: list[tuple[int | None, str]]  # (page number or None, text)
    ocr_pages: int = 0  # Pages read with OCR because they were scans or pictures.
    ocr_skipped: int = 0  # Scanned pages left unread because of MAX_OCR_PAGES.


def _ocr_image(image) -> str:
    """Read the text in a PIL image, top to bottom, one line per row of text."""
    import numpy as np

    image = image.convert("RGB")
    if max(image.size) > OCR_MAX_SIDE:
        image.thumbnail((OCR_MAX_SIDE, OCR_MAX_SIDE))
    with _ocr_lock:
        found, _ = _get_ocr()(np.array(image))
    lines = [(box[0][1], text) for box, text, score in (found or []) if float(score) >= OCR_MIN_CONFIDENCE]
    return "\n".join(text for _, text in sorted(lines, key=lambda line: line[0]))


def _parse_pdf(data: bytes) -> Parsed:
    """Text per page. Pages with no real text are rendered and read with OCR."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [(number, page.extract_text() or "") for number, page in enumerate(reader.pages, start=1)]
    scanned = [i for i, (_, text) in enumerate(pages) if len(text.strip()) < MIN_PAGE_CHARS]
    if not scanned:
        return Parsed(pages)

    import pypdfium2 as pdfium

    # OCR is slow, so read the first MAX_OCR_PAGES scans and say how many were left out.
    to_read = scanned[:MAX_OCR_PAGES]
    pdf = pdfium.PdfDocument(data)
    for i in to_read:
        pages[i] = (pages[i][0], _ocr_image(pdf[i].render(scale=OCR_SCALE).to_pil()))
    return Parsed(pages, ocr_pages=len(to_read), ocr_skipped=len(scanned) - len(to_read))


def _parse_image(data: bytes) -> Parsed:
    from PIL import Image

    return Parsed([(None, _ocr_image(Image.open(io.BytesIO(data))))], ocr_pages=1)


def _parse_docx(data: bytes) -> Parsed:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    # Word files have no fixed pages, so the whole text is one "page".
    return Parsed([(None, "\n".join(parts))])


def _parse_text(data: bytes) -> Parsed:
    try:
        return Parsed([(None, data.decode("utf-8-sig"))])
    except UnicodeDecodeError:
        return Parsed([(None, data.decode("latin-1"))])


def parse_file(filename: str, data: bytes) -> Parsed:
    """Read a file into text, page by page."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_TYPES:
        raise DocumentError(f"Unsupported file type. Use one of: {', '.join(sorted(SUPPORTED_TYPES))}.")
    try:
        if suffix == ".pdf":
            parsed = _parse_pdf(data)
        elif suffix == ".docx":
            parsed = _parse_docx(data)
        elif suffix in IMAGE_TYPES:
            parsed = _parse_image(data)
        else:
            parsed = _parse_text(data)
    except Exception as exc:
        raise DocumentError(f"Could not read the file: {exc}") from exc

    parsed.pages = [(number, text.strip()) for number, text in parsed.pages if text.strip()]
    if not parsed.pages:
        what = "image" if suffix in IMAGE_TYPES else "file"
        hint = " even after OCR" if parsed.ocr_pages else ""
        raise DocumentError(f"No text found in the {what}{hint}.")
    return parsed


# ---------- Embedding model and vector store (created once, on first use) ----------

_lock = threading.Lock()
_embedder = None
_collection = None
_ocr_lock = threading.Lock()  # The OCR engine reads one image at a time.
_ocr = None


def _get_embedder():
    global _embedder
    with _lock:
        if _embedder is None:
            from fastembed import TextEmbedding

            # The first call downloads the model (about 100 MB) into data/models.
            _embedder = TextEmbedding(EMBEDDING_MODEL, cache_dir=str(DATA_DIR / "models"))
        return _embedder


def _get_ocr():
    global _ocr
    if _ocr is None:
        # Small ONNX models that ship inside the package, so there is nothing to download.
        from rapidocr_onnxruntime import RapidOCR

        _ocr = RapidOCR()
    return _ocr


def _get_collection():
    global _collection
    with _lock:
        if _collection is None:
            import chromadb
            from chromadb.config import Settings

            client = chromadb.PersistentClient(
                path=str(DATA_DIR / "chroma"), settings=Settings(anonymized_telemetry=False)
            )
            # We pass our own vectors, so Chroma's built-in embedding function is off.
            _collection = client.get_or_create_collection(
                "documents",
                embedding_function=None,
                configuration={"hnsw": {"space": "cosine"}},
                metadata={"embedding_model": EMBEDDING_MODEL},
            )
            built_with = (_collection.metadata or {}).get("embedding_model")
            if built_with != EMBEDDING_MODEL:
                _collection = None
                raise RuntimeError(
                    f"The document index was built with '{built_with}', but WORLDBEE_EMBEDDING_MODEL "
                    f"is '{EMBEDDING_MODEL}'. Delete {DATA_DIR / 'chroma'} and upload your files again."
                )
        return _collection


# ---------- Keyword index (SQLite FTS5, next to the vector index) ----------

_keyword_lock = threading.Lock()
_keyword_db = None


def _get_keywords() -> sqlite3.Connection:
    """The keyword index. Use it while holding _keyword_lock."""
    global _keyword_db
    if _keyword_db is None:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        db = sqlite3.connect(DATA_DIR / "keywords.sqlite3", check_same_thread=False)
        # Porter stemming, so "trained" matches "training". Only the text is searched.
        db.execute(
            "CREATE VIRTUAL TABLE IF NOT EXISTS chunks USING fts5("
            "text, chunk_id UNINDEXED, doc_id UNINDEXED, tokenize='porter unicode61')"
        )
        # Documents whose chunks are in the index, so older files can be added on first search.
        db.execute("CREATE TABLE IF NOT EXISTS indexed (doc_id TEXT PRIMARY KEY)")
        db.commit()
        _keyword_db = db
    return _keyword_db


def _index_keywords(doc_id: str, ids: list[str], texts: list[str]) -> None:
    with _keyword_lock:
        db = _get_keywords()
        with db:
            db.execute("DELETE FROM chunks WHERE doc_id = ?", (doc_id,))
            db.executemany(
                "INSERT INTO chunks (text, chunk_id, doc_id) VALUES (?, ?, ?)",
                [(text, chunk_id, doc_id) for chunk_id, text in zip(ids, texts)],
            )
            db.execute("INSERT OR IGNORE INTO indexed (doc_id) VALUES (?)", (doc_id,))


def _backfill_keywords(doc_ids: list[str]) -> None:
    """Add files indexed before the keyword index existed, from the text Chroma already holds."""
    with _keyword_lock:
        placeholders = ",".join("?" * len(doc_ids))
        done = {row[0] for row in _get_keywords().execute(
            f"SELECT doc_id FROM indexed WHERE doc_id IN ({placeholders})", doc_ids
        )}
    for doc_id in doc_ids:
        if doc_id not in done:
            stored = _get_collection().get(where={"doc_id": doc_id}, include=["documents"])
            _index_keywords(doc_id, stored["ids"], stored["documents"])


def _keyword_query(query: str) -> str | None:
    """The question as an FTS5 query: any of its words, each quoted so symbols can't break the syntax."""
    words = re.findall(r"\w+", query.lower())
    # PDFs often lose subscripts, so "d_model" in a question is "dmodel" in the text. Try both.
    words += [word.replace("_", "") for word in words if "_" in word]
    return " OR ".join(f'"{word}"' for word in dict.fromkeys(words)) or None


def _keyword_search(query: str, doc_ids: list[str], n: int) -> list[str]:
    """Chunk ids ranked by BM25, best first."""
    match = _keyword_query(query)
    if not match:
        return []
    placeholders = ",".join("?" * len(doc_ids))
    with _keyword_lock:
        rows = _get_keywords().execute(
            f"SELECT chunk_id FROM chunks WHERE chunks MATCH ? AND doc_id IN ({placeholders}) "
            "ORDER BY rank LIMIT ?",
            [match, *doc_ids, n],
        ).fetchall()
    return [row[0] for row in rows]


def _delete_keywords(doc_id: str) -> None:
    with _keyword_lock:
        db = _get_keywords()
        with db:
            db.execute("DELETE FROM chunks WHERE doc_id = ?", (doc_id,))
            db.execute("DELETE FROM indexed WHERE doc_id = ?", (doc_id,))


# ---------- Document list (a small JSON file next to the index) ----------

_registry_lock = threading.Lock()
_REGISTRY = DATA_DIR / "documents.json"


def _read_registry() -> dict[str, dict]:
    try:
        return json.loads(_REGISTRY.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return {}


def _write_registry(registry: dict[str, dict]) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    tmp = _REGISTRY.with_suffix(".tmp")
    tmp.write_text(json.dumps(registry, indent=2), encoding="utf-8")
    tmp.replace(_REGISTRY)


def list_documents() -> list[dict]:
    with _registry_lock:
        return list(_read_registry().values())


def get_documents(doc_ids: list[str]) -> list[dict]:
    """The known documents among doc_ids. Unknown ids are skipped."""
    with _registry_lock:
        registry = _read_registry()
    return [registry[i] for i in doc_ids if i in registry]


# ---------- Public operations ----------


def _find_by_hash(digest: str) -> dict | None:
    with _registry_lock:
        return next((d for d in _read_registry().values() if d.get("sha256") == digest), None)


def _store(filename: str, parsed: Parsed, digest: str, size: int, source_url: str | None = None) -> dict:
    """Chunk, embed and store parsed text, and add it to the library. Returns its record."""
    pages = parsed.pages

    # Split each page on its own, so every chunk knows which page it came from.
    splitter = RecursiveCharacterTextSplitter(chunk_size=CHUNK_SIZE, chunk_overlap=CHUNK_OVERLAP)
    chunks: list[tuple[int | None, str]] = [
        (number, piece) for number, text in pages for piece in splitter.split_text(text)
    ]

    doc_id = uuid.uuid4().hex
    texts = [text for _, text in chunks]
    embeddings = list(_get_embedder().embed(texts, batch_size=32))
    metadatas = []
    for index, (number, _) in enumerate(chunks):
        meta = {"doc_id": doc_id, "filename": filename, "chunk": index}
        if number is not None:  # Chroma metadata can't hold None.
            meta["page"] = number
        if source_url:
            meta["source_url"] = source_url
        metadatas.append(meta)

    ids = [f"{doc_id}:{i}" for i in range(len(chunks))]
    _get_collection().add(ids=ids, embeddings=embeddings, documents=texts, metadatas=metadatas)
    try:
        _index_keywords(doc_id, ids, texts)
    except sqlite3.Error:
        pass  # Search adds it later (see _backfill_keywords) and uses vectors meanwhile.

    record = {
        "id": doc_id,
        "filename": filename,
        "size": size,
        "sha256": digest,
        # Only PDFs have real pages.
        "pages": len(pages) if Path(filename).suffix.lower() == ".pdf" else None,
        # Pages that were scans or pictures read with OCR, and scans left unread (0 for ordinary files).
        "ocr_pages": parsed.ocr_pages,
        "ocr_skipped": parsed.ocr_skipped,
        # Set when the file came from a web address.
        "source_url": source_url,
        "chunks": len(chunks),
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }
    with _registry_lock:
        registry = _read_registry()
        registry[doc_id] = record
        _write_registry(registry)
    return record


def add_document(filename: str, data: bytes, source_url: str | None = None) -> dict:
    """Parse, chunk, embed and store one file. Returns its record.

    The same file uploaded twice is only indexed once: the library keeps the
    first copy and returns its record.
    """
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError(f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    if not data:
        raise DocumentError("The file is empty.")

    digest = hashlib.sha256(data).hexdigest()
    existing = _find_by_hash(digest)
    if existing:
        return existing
    return _store(filename, parse_file(filename, data), digest, len(data), source_url)


# File types a link can point at, by what the site says it is.
_LINKED_FILE_EXTENSIONS = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/webp": ".webp",
    "image/bmp": ".bmp",
    "image/tiff": ".tiff",
}


def add_url(raw_url: str) -> dict:
    """Fetch a web address and index what is there: a web page, or a PDF, Word, text or image file."""
    try:
        fetched = web.fetch(raw_url, MAX_FILE_BYTES)
        parts = urlparse(fetched.url)
        host = parts.hostname.removeprefix("www.")
        if fetched.content_type in ("text/html", "application/xhtml+xml"):
            title, text = web.extract_page(fetched.data, fetched.url)
            text_bytes = text.encode()
            digest = hashlib.sha256(text_bytes).hexdigest()
            existing = _find_by_hash(digest)
            if existing:
                return existing
            name = " ".join((title or f"{host}{parts.path.rstrip('/')}").split())[:90]
            return _store(name, Parsed([(None, text)]), digest, len(text_bytes), fetched.url)
        suffix = _LINKED_FILE_EXTENSIONS.get(fetched.content_type) or Path(parts.path).suffix.lower()
        if suffix not in SUPPORTED_TYPES:
            raise DocumentError(
                f"That link ({fetched.content_type or 'unknown type'}) can't be indexed. "
                "Use a web page, PDF, Word, text or image link."
            )
        stem = Path(parts.path).stem or host
        return add_document(f"{stem}{suffix}", fetched.data, source_url=fetched.url)
    except web.FetchError as exc:
        raise DocumentError(str(exc)) from None


def delete_document(doc_id: str) -> bool:
    """Remove a document and its chunks. Returns False if it didn't exist."""
    with _registry_lock:
        registry = _read_registry()
        if doc_id not in registry:
            return False
        _get_collection().delete(where={"doc_id": doc_id})
        _delete_keywords(doc_id)
        del registry[doc_id]
        _write_registry(registry)
    return True


def _vector_search(query: str, doc_ids: list[str], n: int) -> list[str]:
    """Chunk ids ranked by meaning, closest first."""
    query_vector = next(iter(_get_embedder().query_embed(query)))
    result = _get_collection().query(
        query_embeddings=[query_vector], n_results=n, where={"doc_id": {"$in": doc_ids}}, include=[]
    )
    return result["ids"][0]


def search(query: str, doc_ids: list[str], k: int = TOP_K, mode: str = "hybrid") -> list[dict]:
    """The k best chunks for the query, from the given documents only.

    mode is "hybrid" (the default), or "vector" or "keyword" for one side alone,
    which the retrieval evaluation uses for comparison.
    """
    if not doc_ids:
        return []
    ranked_lists = []
    if mode in ("hybrid", "vector"):
        ranked_lists.append(_vector_search(query, doc_ids, CANDIDATES if mode == "hybrid" else k))
    if mode in ("hybrid", "keyword"):
        try:
            _backfill_keywords(doc_ids)
            ranked_lists.append(_keyword_search(query, doc_ids, CANDIDATES if mode == "hybrid" else k))
        except sqlite3.Error:
            if mode == "keyword":
                raise  # In hybrid mode the vector results still stand on their own.

    # Reciprocal rank fusion: a chunk near the top of either list, or fairly high in both, wins.
    fused: dict[str, float] = {}
    for ranked in ranked_lists:
        for rank, chunk_id in enumerate(ranked, start=1):
            fused[chunk_id] = fused.get(chunk_id, 0.0) + 1 / (RRF_K + rank)
    best = sorted(fused, key=fused.get, reverse=True)[:k]
    if not best:
        return []

    stored = _get_collection().get(ids=best, include=["documents", "metadatas"])
    by_id = dict(zip(stored["ids"], zip(stored["documents"], stored["metadatas"])))
    return [
        {
            "text": text,
            "filename": meta["filename"],
            "page": meta.get("page"),
            "url": meta.get("source_url"),
        }
        for text, meta in (by_id[chunk_id] for chunk_id in best if chunk_id in by_id)
    ]
